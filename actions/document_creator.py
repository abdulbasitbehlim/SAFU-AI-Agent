"""Create common user documents at real Windows user-folder locations."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from core.user_folders import user_folder
from actions.file_controller import _mutation_allowed, _resolve_path
from core.undo import push_undo


def _target_path(location: str, filename: str) -> Path:
    loc = (location or "desktop").strip()
    base = _resolve_path(loc)
    name = (filename or "Safu_document.txt").strip()
    return base / name


def _undo_created(path: Path):
    def _fn():
        if path.exists() and path.is_file():
            path.unlink()
            return f"Removed '{path.name}'."
        return f"'{path.name}' is already gone."
    return _fn


def _json_data(raw: str, default):
    if not raw:
        return default
    obj = json.loads(raw)
    return obj


def document_creator(parameters=None, response=None, player=None, session_memory=None) -> str:
    p = parameters or {}
    filename = (p.get("filename") or "Safu_document.txt").strip()
    location = (p.get("location") or "desktop").strip()
    content = p.get("content") or ""
    data_json = p.get("data_json") or ""
    title = (p.get("title") or Path(filename).stem or "Document").strip()

    target = _target_path(location, filename)
    # Never overwrite an existing user document implicitly. Choose a duplicate-safe
    # name instead; the user can explicitly use file_controller.write if they intend
    # to replace existing contents.
    if target.exists():
        stem, suffix = target.stem, target.suffix
        i = 1
        while True:
            candidate = target.with_name(f"{stem} ({i}){suffix}")
            if not candidate.exists():
                target = candidate
                break
            i += 1
    ok, why = _mutation_allowed(target)
    if not ok:
        return why
    target.parent.mkdir(parents=True, exist_ok=True)
    ext = target.suffix.lower()

    try:
        if ext in {".txt", ".md", ".json", ".py", ".csv"}:
            if ext == ".json" and data_json:
                obj = _json_data(data_json, {})
                target.write_text(json.dumps(obj, indent=2, ensure_ascii=False), encoding="utf-8")
            elif ext == ".csv" and data_json:
                rows = _json_data(data_json, [])
                if not isinstance(rows, list):
                    return "CSV data_json must be a JSON list."
                with target.open("w", newline="", encoding="utf-8-sig") as f:
                    if rows and isinstance(rows[0], dict):
                        fields = list(dict.fromkeys(k for r in rows if isinstance(r, dict) for k in r.keys()))
                        w = csv.DictWriter(f, fieldnames=fields)
                        w.writeheader(); w.writerows(r for r in rows if isinstance(r, dict))
                    else:
                        w = csv.writer(f)
                        for row in rows:
                            w.writerow(row if isinstance(row, list) else [row])
            else:
                target.write_text(str(content), encoding="utf-8")

        elif ext == ".docx":
            from docx import Document
            doc = Document()
            if title:
                doc.add_heading(title, level=1)
            for block in str(content).split("\n\n"):
                block = block.strip()
                if block:
                    doc.add_paragraph(block)
            doc.save(str(target))

        elif ext == ".xlsx":
            from openpyxl import Workbook
            wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
            rows = _json_data(data_json, []) if data_json else []
            if rows and isinstance(rows, list) and isinstance(rows[0], dict):
                fields = list(dict.fromkeys(k for r in rows if isinstance(r, dict) for k in r.keys()))
                ws.append(fields)
                for r in rows:
                    if isinstance(r, dict): ws.append([r.get(k, "") for k in fields])
            elif rows and isinstance(rows, list):
                for r in rows: ws.append(r if isinstance(r, list) else [r])
            elif content:
                for line in str(content).splitlines(): ws.append([line])
            wb.save(str(target))

        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation()
            slides = _json_data(data_json, []) if data_json else []
            if not slides:
                slides = [{"title": title, "body": content}]
            for item in slides:
                item = item if isinstance(item, dict) else {"title": "", "body": str(item)}
                layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(layout)
                slide.shapes.title.text = str(item.get("title", ""))
                slide.placeholders[1].text = str(item.get("body", ""))
            prs.save(str(target))

        else:
            return "Supported document types: txt, md, json, csv, docx, xlsx, pptx."

        push_undo(f"created {target.name}", _undo_created(target))
        if player:
            player.write_log(f"[document] created {target}")
        return f"Created: {target}"
    except Exception as e:
        return f"Could not create document: {e}"


TOOL = {
    "name": "document_creator",
    "description": (
        "Creates real files/documents on Desktop or another local folder. Supports TXT, Markdown, JSON, CSV, Word DOCX, Excel XLSX and PowerPoint PPTX. "
        "Use this instead of merely describing a document when the user asks to make/save/create a file."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "location": {"type": "STRING", "description": "desktop, documents, downloads, D:/path, or another local folder"},
            "filename": {"type": "STRING", "description": "Filename including extension, e.g. notes.docx"},
            "title": {"type": "STRING", "description": "Optional document title"},
            "content": {"type": "STRING", "description": "Text content"},
            "data_json": {"type": "STRING", "description": "Optional JSON string for table rows/slides; list of dicts works for XLSX/CSV, list of {title,body} for PPTX"},
        },
        "required": ["filename"],
    },
    "handler": document_creator,
}
